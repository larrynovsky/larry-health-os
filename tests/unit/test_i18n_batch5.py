"""K10: frozen copy, tenant language and read-time views, without database/network I/O.

The staging file is a review artifact. Production uses the canonical dictionaries;
merge the two maps before enabling the code. The tests also work after that merge.
"""
from __future__ import annotations

import ast
import hashlib
import json
import re
import sys
from copy import deepcopy
from pathlib import Path
from string import Formatter
from types import SimpleNamespace

import pytest
import yaml

pytestmark = pytest.mark.unit
ROOT = Path(__file__).resolve().parents[2]
_K10_KEYS = (
    "cpic.drug.warfarin",
    "cpic.drug.celecoxib",
    "cpic.drug.phenytoin",
    "cpic.drug.ibuprofen",
    "cpic.drug.clopidogrel",
    "cpic.drug.omeprazole",
    "cpic.drug.escitalopram",
    "cpic.drug.voriconazole",
    "cpic.drug.codeine",
    "cpic.drug.tamoxifen",
    "cpic.drug.metoprolol",
    "cpic.drug.fluorouracil",
    "cpic.drug.azathioprine",
    "cpic.drug.irinotecan",
    "cpic.drug.atazanavir",
    "cpic.drug.simvastatin",
    "cpic.drug.atorvastatin",
    "cpic.drug.rosuvastatin",
    "cpic.drug.pravastatin",
    "cpic.drug.abacavir",
    "cpic.drug.flucloxacillin",
    "cpic.implication.cyp2c9.poor_metabolizer",
    "cpic.implication.cyp2c9.intermediate_metabolizer",
    "cpic.implication.cyp2c19.poor_metabolizer",
    "cpic.implication.cyp2c19.intermediate_metabolizer",
    "cpic.implication.cyp2c19.ultrarapid_metabolizer",
    "cpic.implication.cyp2c19.rapid_metabolizer",
    "cpic.implication.cyp2d6.poor_metabolizer",
    "cpic.implication.cyp2d6.ultrarapid_metabolizer",
    "cpic.implication.cyp2d6.intermediate_metabolizer",
    "cpic.implication.slco1b1.poor_function",
    "cpic.implication.slco1b1.decreased_function",
    "cpic.implication.dpyd.poor_metabolizer",
    "cpic.implication.dpyd.intermediate_metabolizer",
    "cpic.implication.tpmt.poor_metabolizer",
    "cpic.implication.tpmt.intermediate_metabolizer",
    "cpic.implication.ugt1a1.poor_metabolizer",
    "cpic.implication.ugt1a1.intermediate_metabolizer",
    "cpic.implication.hla_b.hla_b_57_01_carrier",
    "cpic.phenotype.undetermined",
    "cpic.advice.no_genotype",
    "cpic.advice.no_recommendation",
    "traits.notes.lactase_missing",
    "traits.notes.lactase_persistent",
    "traits.notes.lactase_partial",
    "traits.notes.lactase_nonpersistent",
    "traits.notes.eye_missing",
    "traits.notes.eye_light",
    "traits.notes.eye_mixed",
    "traits.notes.eye_dark",
    "traits.notes.apoe_missing",
    "traits.notes.apoe_e3_e3",
    "traits.notes.apoe_e2_e3",
    "traits.notes.apoe_e2_e2",
    "traits.notes.apoe_e3_e4",
    "traits.notes.apoe_e4_e4",
    "traits.notes.apoe_phase_uncertain",
    "traits.notes.apoe_unusual",
    "wellness.notes.mthfr_missing",
    "wellness.notes.mthfr_normal",
    "wellness.notes.mthfr_677_het",
    "wellness.notes.mthfr_1298_het",
    "wellness.notes.mthfr_677_hom",
    "wellness.notes.mthfr_compound",
    "wellness.notes.mthfr_1298_hom",
    "wellness.notes.mthfr_severe",
    "wellness.notes.mthfr_rare",
    "wellness.notes.comt_missing",
    "wellness.notes.comt_low",
    "wellness.notes.comt_intermediate",
    "wellness.notes.comt_high",
    "wellness.notes.fto_missing",
    "wellness.notes.fto_elevated",
    "wellness.notes.fto_modest",
    "wellness.notes.fto_neutral",
    "wellness.notes.hfe_missing",
    "wellness.notes.hfe_homozygous",
    "wellness.notes.hfe_heterozygous",
    "wellness.notes.hfe_absent",
    "longitudinal.label.sleep_total",
    "longitudinal.label.sleep_deep",
    "longitudinal.label.sleep_rem",
    "longitudinal.label.hrv",
    "longitudinal.label.active_kcal",
    "longitudinal.phase.before_procedure",
    "longitudinal.sheet.yearly",
    "longitudinal.column.year",
    "longitudinal.legend.row_colours",
    "longitudinal.sheet.phases",
    "longitudinal.column.phase",
    "longitudinal.column.type",
    "longitudinal.column.start",
    "longitudinal.column.end",
    "longitudinal.column.days",
    "longitudinal.sheet.correlations",
    "longitudinal.heading.spearman",
    "longitudinal.column.metric_a",
    "longitudinal.column.metric_b",
    "longitudinal.column.significant",
    "longitudinal.column.gate",
    "longitudinal.gate.passed",
    "longitudinal.gate.phantom",
    "longitudinal.heading.lagged",
    "longitudinal.column.predictor",
    "longitudinal.column.target",
    "longitudinal.column.lag",
    "longitudinal.heading.labs",
    "longitudinal.column.lab",
    "longitudinal.column.metric",
    "longitudinal.column.strong",
    "longitudinal.sheet.recovery",
    "longitudinal.column.baseline",
    "longitudinal.column.current",
    "longitudinal.column.percent_baseline",
    "longitudinal.column.trend",
    "longitudinal.links.header",
    "longitudinal.links.none",
    "longitudinal.links.previous_gone",
    "longitudinal.links.confirmed",
    "longitudinal.links.new",
    "longitudinal.links.positive",
    "longitudinal.links.negative",
    "longitudinal.links.gone",
    "longitudinal.links.caution",
)
_K10_RU_SHA256 = "b762381859c16cc249f1ec831e4f5e22e2bb89e30a636da40ce6f76df0941eda"
_K10_EN_SHA256 = "b2b9e82a47e05afd15549c7cf9ef26f288f6e44a0d233ff4a9d0ef629f838987"


def _k10_tables():
    tables = {lang: yaml.safe_load((ROOT / f"methodology/i18n/{lang}.yaml").read_text())
              for lang in ("ru", "en")}
    staged = ROOT / "plans/K10_NEW_KEYS_2026-09-29.yaml"
    if staged.exists():
        new = yaml.safe_load(staged.read_text())
        assert set(new) == {"ru", "en"}
        for lang in tables:
            assert set(new[lang]) == set(_K10_KEYS)
            for key in set(tables[lang]) & set(new[lang]):
                assert tables[lang][key] == new[lang][key], key
            tables[lang].update(new[lang])
    return tables


def _k10_bind(monkeypatch):
    import i18n
    tables = _k10_tables()
    monkeypatch.setattr(i18n, "_table", tables.__getitem__)
    # Real lang_of resolves the current tenant profile every time. No owner profile,
    # database adapter, paths module, secret store or actual profile is imported.
    profile = {"identity.language": "ru"}
    monkeypatch.setitem(sys.modules, "profile_db", SimpleNamespace(
        get_patient_profile=lambda: dict(profile)))
    return i18n, tables, profile


def _k10_digest(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True,
                                     separators=(",", ":")).encode()).hexdigest()


def _k10_longitudinal():
    """Execute the actual rendering functions, excluding import-time tenant paths/DB."""
    import i18n
    import openpyxl
    import pandas as pd
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    tree = ast.parse((ROOT / "longitudinal_analysis.py").read_text())
    names = {"_longitudinal_label", "links_note", "_header_row", "_auto_width",
             "write_sheet_yearly", "write_sheet_phases", "write_sheet_correlations",
             "write_sheet_recovery"}
    constants = {"_DAILY_LABELS", "_LONGITUDINAL_LEGACY_LABELS",
                 "_LONGITUDINAL_LABEL_KEYS", "PHASE_COLORS"}
    nodes = [n for n in tree.body if
             isinstance(n, ast.FunctionDef) and n.name in names or
             isinstance(n, ast.Assign) and any(isinstance(t, ast.Name) and t.id in constants
                                               for t in n.targets)]
    assert {n.name for n in nodes if isinstance(n, ast.FunctionDef)} == names
    ns = dict(i18n=i18n, openpyxl=openpyxl, pd=pd, Font=Font, PatternFill=PatternFill,
              Alignment=Alignment, get_column_letter=get_column_letter)
    code = ast.fix_missing_locations(ast.Module(body=[
        ast.ImportFrom(module="__future__", names=[ast.alias(name="annotations")], level=0),
        *nodes], type_ignores=[]))
    exec(compile(code, "longitudinal_analysis.py", "exec"), ns)
    # The production membership filter does not change the labels under test.
    ns["DAILY_METRICS"] = ns["_DAILY_LABELS"]
    return SimpleNamespace(**ns)


def test_k10_dictionary_snapshot_placeholders_and_vocabulary():
    tables = _k10_tables()
    assert _k10_digest({k: tables["ru"][k] for k in _K10_KEYS}) == _K10_RU_SHA256
    assert _k10_digest({k: tables["en"][k] for k in _K10_KEYS}) == _K10_EN_SHA256
    # Reuse the exact vocabulary oracle, without importing its I/O fixtures.
    source = ast.parse((ROOT / "tests/unit/test_i18n.py").read_text())
    check = next(n for n in source.body if isinstance(n, ast.FunctionDef) and
                 n.name == "test_person_copy_has_no_internal_vocabulary_or_formal_address")
    forbidden = next(n.value.args[0].value for n in check.body if isinstance(n, ast.Assign)
                     and any(isinstance(t, ast.Name) and t.id == "forbidden" for t in n.targets))
    for key in _K10_KEYS:
        ru, en = tables["ru"][key], tables["en"][key]
        assert not re.search(forbidden, re.sub(r"\{[^}]*\}", "", ru), re.I), key
        assert not re.search("[А-Яа-яЁё]", en), key
        assert {(f, s, c) for _, f, s, c in Formatter().parse(ru) if f is not None} == {
            (f, s, c) for _, f, s, c in Formatter().parse(en) if f is not None}, key


def test_k10_stored_notes_translate_without_mutation_or_recalling_genotypes(monkeypatch):
    i18n, tables, profile = _k10_bind(monkeypatch)
    import traits_pipeline as traits
    import wellness_pipeline as wellness

    for module, keys, legacy, render in (
        (traits, traits._TRAIT_NOTE_KEYS, traits._TRAIT_LEGACY_NOTES, traits.trait_notes_for_person),
        (wellness, wellness._WELLNESS_NOTE_KEYS, wellness._WELLNESS_LEGACY_NOTES,
         wellness.wellness_notes_for_person),
    ):
        # Rows intentionally have no genotype/phenotype fields: the persisted note is
        # the source, and calling the genetic pipeline again would corrupt its meaning.
        rows = [{"notes": legacy.get(key, tables["ru"][key]).format(missing="['demo_marker']")}
                for key in keys]
        rows += [{"notes": "Invented custom note"}, {"notes": None}, {"notes": ""}]
        original = deepcopy(rows)
        for lang in ("en", "ru", "en"):
            profile["identity.language"] = lang
            actual = render(rows)
            assert [r["notes"] for r in actual[:-3]] == [
                tables[lang][key].format(missing="['demo_marker']") for key in keys]
            assert actual[-3:] == original[-3:]
            assert rows == original
            assert all(a is not b for a, b in zip(actual, rows))
        # Explicit report language beats the current process profile.
        assert render(rows, "ru")[0]["notes"] == tables["ru"][keys[0]]


def test_k10_pipeline_writers_keep_russian_storage_under_english_profile(monkeypatch):
    _, tables, profile = _k10_bind(monkeypatch)
    profile["identity.language"] = "en"
    import traits_pipeline as traits
    import wellness_pipeline as wellness

    cases = [
        (traits._call_lactase, "rs4988235", ("AA", "AG", "GG", None),
         ("lactase_persistent", "lactase_partial", "lactase_nonpersistent", "lactase_missing")),
        (traits._call_eye_color, "rs12913832", ("GG", "AG", "AA", None),
         ("eye_light", "eye_mixed", "eye_dark", "eye_missing")),
        (wellness._call_comt, "rs4680", ("AA", "GA", "GG", None),
         ("comt_low", "comt_intermediate", "comt_high", "comt_missing")),
        (wellness._call_fto, "rs9939609", ("AA", "AT", "TT", None),
         ("fto_elevated", "fto_modest", "fto_neutral", "fto_missing")),
        (wellness._call_hfe, "rs1799945", ("GG", "CG", "CC", None),
         ("hfe_homozygous", "hfe_heterozygous", "hfe_absent", "hfe_missing")),
    ]
    for caller, marker, genotypes, suffixes in cases:
        prefix = "traits" if caller.__module__ == "traits_pipeline" else "wellness"
        legacy = traits._TRAIT_LEGACY_NOTES if prefix == "traits" else wellness._WELLNESS_LEGACY_NOTES
        for genotype, suffix in zip(genotypes, suffixes):
            key = f"{prefix}.notes.{suffix}"
            assert caller({marker: genotype}, {}).notes == legacy.get(key, tables["ru"][key])

    for g1, g2, suffix in (
        ("GG", "TT", "normal"), ("GA", "TT", "677_het"), ("GG", "TG", "1298_het"),
        ("AA", "TT", "677_hom"), ("GA", "TG", "compound"), ("GG", "GG", "1298_hom"),
        ("AA", "TG", "severe"), ("AA", "GG", "rare"), (None, None, "missing"),
    ):
        key = "wellness.notes.mthfr_" + suffix
        assert wellness._call_mthfr({"rs1801133": g1, "rs1801131": g2}, {}).notes == (
            wellness._WELLNESS_LEGACY_NOTES.get(key, tables["ru"][key]))

    for g1, g2, suffix in (
        ("CC", "TT", "e3_e3"), ("CT", "TT", "e2_e3"), ("TT", "TT", "e2_e2"),
        ("CC", "CT", "e3_e4"), ("CC", "CC", "e4_e4"), ("CT", "CT", "phase_uncertain"),
        ("TT", "CC", "unusual"), (None, None, "missing"),
    ):
        key = "traits.notes.apoe_" + suffix
        assert traits._call_apoe({"rs7412": g1, "rs429358": g2}, {}).notes == (
            traits._TRAIT_LEGACY_NOTES.get(key, tables["ru"][key]).format(
                missing=["rs7412", "rs429358"]))


def test_k10_cpic_old_catalog_and_implications_follow_tenant_language(monkeypatch):
    _, tables, profile = _k10_bind(monkeypatch)
    import cpic_reference_db as cpic
    # Simulate the low-level read boundary; no SQLite connection or seed is used.
    rows = [{"drug_id": d["drug_id"], "drug_display": tables["ru"][d["drug_display_key"]],
             "drug_search": d["drug_search"], "gene": d["gene"]} for d in cpic._DRUGS]
    original = deepcopy(rows)
    monkeypatch.setattr(cpic, "drug_catalog", lambda conn: rows)
    monkeypatch.setattr(cpic, "drug_risk_for", lambda *args: None)
    monkeypatch.setattr(cpic, "get_gene_implications", lambda gene, pheno, conn: [
        tables["ru"][key] for g, pm, key in cpic._GENE_IMPLICATIONS
        if g == gene and pm.lower() in pheno.lower()])
    for lang in ("en", "ru", "en"):
        profile["identity.language"] = lang
        no_genome = cpic.build_drug_interactions({}, conn=object())
        assert [r["drug"] for r in no_genome] == [
            tables[lang][d["drug_display_key"]] for d in cpic._DRUGS]
        assert all(r["phenotype_context"] == tables[lang]["cpic.phenotype.undetermined"]
                   and r["recommendation"] == tables[lang]["cpic.advice.no_genotype"]
                   and r["risk_level"] == "unknown" for r in no_genome)
        for gene, phenotype, key in cpic._GENE_IMPLICATIONS:
            output = cpic.build_drug_interactions({gene: phenotype}, conn=object())
            # Preserve the existing substring semantics (Rapid also matches Ultrarapid).
            matched = [k for g, pm, k in cpic._GENE_IMPLICATIONS
                       if g == gene and pm.lower() in phenotype.lower()]
            markers = " ".join(tables["ru"][k] for k in matched)
            expected_risk = "critical" if "🚨" in markers else (
                "warning" if "⚠️" in markers or "⚡" in markers else "ok")
            assert all(r["recommendation"] == " ".join(tables[lang][k] for k in matched) and
                       r["risk_level"] == expected_risk for r in output if r["gene"] == gene)
        assert rows == original
    output = cpic.build_drug_interactions({"CYP2C9": "invented class"}, conn=object())
    assert all(r["recommendation"] == tables["en"]["cpic.advice.no_recommendation"]
               for r in output if r["gene"] == "CYP2C9")
    monkeypatch.setattr(cpic, "drug_risk_for", lambda *args: {
        "risk_level": "warning", "recommendation": "Invented custom recommendation"})
    output = cpic.build_drug_interactions({"CYP2C9": "invented class"}, conn=object())
    assert all(r["recommendation"] == "Invented custom recommendation"
               for r in output if r["gene"] == "CYP2C9")
    rows[0]["drug_display"] = "Invented custom name"
    assert cpic.build_drug_interactions({}, conn=object())[0]["drug"] == "Invented custom name"


def test_k10_seed_keeps_original_catalog_bytes(monkeypatch):
    _, _, profile = _k10_bind(monkeypatch)
    profile["identity.language"] = "en"
    import cpic_reference_db as cpic
    # A recording object only: no database, files, snapshot or private overlay opened.
    writes = []
    conn = SimpleNamespace(execute=lambda sql, args=(): writes.append((sql, args)),
                           executescript=lambda sql: None, commit=lambda: None)
    monkeypatch.setattr(cpic, "seed_version", lambda: "demo-version")
    monkeypatch.setattr(cpic, "_current_seed_version", lambda conn: None)
    monkeypatch.setattr(cpic, "_snapshot_rows", lambda: [])
    monkeypatch.setattr(cpic, "_read_json", lambda path: {})
    cpic.seed(conn)
    catalog = [dict(zip(("drug_id", "drug_display", "drug_search", "gene"), args[:4]))
               for sql, args in writes if sql.startswith("INSERT INTO cpic_drug_catalog")]
    implications = [args[:3] for sql, args in writes
                    if sql.startswith("INSERT INTO cpic_gene_implication")]
    # Pinned from the pre-K10 source AST, independent of the new keys/translations.
    assert _k10_digest({"_DRUGS": catalog, "_GENE_IMPLICATIONS": implications}) == (
        "be19eb25a50ea5e2ed77bbf36df55e8f89a92e51b4a983701d118435a3c90fa6")


def test_k10_read_views_leave_unknown_implications_unchanged(monkeypatch):
    _, tables, profile = _k10_bind(monkeypatch)
    profile["identity.language"] = "en"
    import cpic_reference_db as cpic
    key = cpic._GENE_IMPLICATIONS[0][2]
    rows = [("Poor Metabolizer", tables["ru"][key]),
            ("Poor Metabolizer", "Invented custom implication"),
            ("Normal Metabolizer", "Unrelated class")]
    query_rows = SimpleNamespace(fetchall=lambda: rows)
    conn = SimpleNamespace(execute=lambda *args: query_rows)
    # Exercise both real readers: the raw view must remain raw.
    assert cpic.get_gene_implications("CYP2C9", "Poor Metabolizer", conn) == [
        tables["ru"][key], "Invented custom implication"]
    assert cpic.cpic_implications_for_person("CYP2C9", "Poor Metabolizer", conn) == [
        tables["en"][key], "Invented custom implication"]


def test_k10_links_note_switches_existing_russian_labels_and_all_branches(monkeypatch):
    _, tables, profile = _k10_bind(monkeypatch)
    la = _k10_longitudinal()
    cur = [("a×b", tables["ru"]["longitudinal.label.sleep_total"],
            tables["ru"]["safety.label.resting_hr"], 0.23),
           ("c×d", tables["ru"]["experiments.metric.steps"], "X", -0.31)]
    original = deepcopy(cur)
    for lang in ("ru", "en", "ru"):
        profile["identity.language"] = lang
        note = la.links_note(["a×b", "gone"], cur, 7, "2040-01-01")
        t = tables[lang]
        assert note == "\n".join([
            t["longitudinal.links.header"].format(when="2040-01-01", m=7),
            t["longitudinal.links.confirmed"],
            f'• {t["longitudinal.label.sleep_total"]} ↔ {t["safety.label.resting_hr"]}: '
            f'r=+0.23 ({t["longitudinal.links.positive"]})',
            f'• {t["experiments.metric.steps"]} ↔ X: r=-0.31 '
            f'({t["longitudinal.links.negative"]}){t["longitudinal.links.new"]}',
            t["longitudinal.links.gone"].format(pairs="gone"), t["longitudinal.links.caution"]])
        for prev in (None, [], ["gone"]):
            assert la.links_note(prev, [], 7, "2040-01-01") == (
                t["longitudinal.links.header"].format(when="2040-01-01", m=7) +
                t["longitudinal.links.none"] + (t["longitudinal.links.previous_gone"].format(
                    pairs="gone") if prev else ""))
        assert tables[lang]["longitudinal.links.new"] not in la.links_note(None, cur, 7, "date")
    assert cur == original


def test_k10_workbook_sheets_and_rows_use_report_language(monkeypatch):
    _, tables, profile = _k10_bind(monkeypatch)
    import openpyxl
    import pandas as pd
    la = _k10_longitudinal()
    # Synthetic values describe test rows, not a person or a medical record.
    corr = pd.DataFrame([dict(metric_a="sleep_total", metric_b="steps", spearman_r=0.3,
                             p_value=0.01, n=30, significant=True, derived=False,
                             gate_pass=True, p_perm=0.01)])
    lagged = pd.DataFrame([dict(predictor="x", target="y", lag_days=1, spearman_r=-0.3,
                               p_value=0.01, n=30, significant=True)])
    labs = pd.DataFrame([dict(lab="DEMO_ANALYTE", metric="steps", spearman_r=0.3,
                             p_value=0.01, n=30, significant=True, strong=True,
                             gate_pass=True, p_perm=0.01)])
    yearly = pd.DataFrame([{"year": 2040, "steps_mean": 1}])
    phases = pd.DataFrame([dict(phase="DEMO", phase_type="baseline", start="2040-01-01",
                               end="2040-01-02", n_days=2, steps_mean=1)])
    recovery = {"steps": dict(pct_of_baseline=100, baseline_mean=1, baseline_p25=1,
                              baseline_p75=1, current_mean=1, trend_dir="→")}
    for lang in ("ru", "en"):
        profile["identity.language"] = "ru" if lang == "en" else "en"
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        la.write_sheet_yearly(wb, yearly, [], lang=lang)
        la.write_sheet_phases(wb, phases, lang=lang)
        la.write_sheet_correlations(wb, corr, lagged, labs, lang=lang)
        la.write_sheet_recovery(wb, recovery, lang=lang)
        assert wb.sheetnames == [tables[lang]["longitudinal.sheet." + suffix]
                                for suffix in ("yearly", "phases", "correlations", "recovery")]
        cells = [c.value for ws in wb for row in ws for c in row if isinstance(c.value, str)]
        assert tables[lang]["longitudinal.label.sleep_total"] in cells
        assert tables[lang]["experiments.metric.steps"] in cells
        assert tables[lang]["longitudinal.gate.passed"] in cells
        if lang == "en":
            assert not any(re.search("[А-Яа-яЁё]", value) for value in cells)
        profile["identity.language"] = lang
        wb_default = openpyxl.Workbook()
        la.write_sheet_recovery(wb_default, recovery)
        assert wb_default.worksheets[-1].title == tables[lang]["longitudinal.sheet.recovery"]


def test_k10_person_entrypoints_resolve_tenant_language_without_owner_lookup():
    entries = {
        "cpic_reference_db.py": {"build_drug_interactions", "cpic_implications_for_person"},
        "traits_pipeline.py": {"_genetic_notes_for_person"},
        "longitudinal_analysis.py": {"links_note", "load_phases", "write_sheet_yearly",
                                     "write_sheet_phases", "write_sheet_correlations",
                                     "write_sheet_recovery"},
    }
    for filename, names in entries.items():
        source = (ROOT / filename).read_text()
        functions = {n.name: n for n in ast.parse(source).body if isinstance(n, ast.FunctionDef)}
        for name in names:
            function = functions[name]
            assert any(isinstance(n, ast.Call) and isinstance(n.func, ast.Attribute) and
                       isinstance(n.func.value, ast.Name) and n.func.value.id == "i18n" and
                       n.func.attr == "lang_of" and not n.args for n in ast.walk(function))
            assert "owner" not in ast.get_source_segment(source, function).lower()


def test_k10_genome_read_path_uses_note_views():
    """Integration oracle: expected RED until the reviewer wires the two read views.

    K10 explicitly forbids edits to dashboard files. Leaving this check unskipped
    prevents green unit renderers from being mistaken for translated person pages.
    """
    tree = ast.parse((ROOT / "dashboard_routers/views.py").read_text())
    genome = next(n for n in tree.body if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef))
                  and n.name == "genome")
    calls = {n.func.id if isinstance(n.func, ast.Name) else n.func.attr
             for n in ast.walk(genome) if isinstance(n, ast.Call)
             and isinstance(n.func, (ast.Name, ast.Attribute))}
    assert {"trait_notes_for_person", "wellness_notes_for_person"} <= calls, (
        "K10 integration pending: pass the stored trait/wellness rows through their person views "
        "in dashboard_routers/views.py::genome; that file is outside this batch's write scope")
